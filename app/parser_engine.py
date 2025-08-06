import pdfplumber
import pytesseract
from PIL import Image
from docx import Document
from openpyxl import load_workbook
import json
import os

def process_document(path, template_name):
    ext = os.path.splitext(path)[-1].lower()
    with open(f"templates/{template_name}.json") as f:
        template = json.load(f)

    if ext == ".pdf":
        return parse_pdf(path, template)
    elif ext == ".docx":
        return parse_docx(path, template)
    elif ext in [".xlsx", ".xls"]:
        return parse_excel(path, template)
    elif ext in [".png", ".jpg", ".jpeg"]:
        return parse_image(path, template)
    else:
        return {"error": "Unsupported file type"}

def parse_pdf(path, template):
    data = {}
    with pdfplumber.open(path) as pdf:
        text = "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
        for field, rule in template.items():
            data[field] = extract_by_rule(text, rule)
    return data

def parse_docx(path, template):
    data = {}
    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs])
    for field, rule in template.items():
        data[field] = extract_by_rule(text, rule)
    return data

def parse_excel(path, template):
    wb = load_workbook(path)
    sheet = wb.active
    data = {}
    for field, rule in template.items():
        cell = sheet[rule]
        data[field] = cell.value
    return data

def parse_image(path, template):
    image = Image.open(path)
    text = pytesseract.image_to_string(image)
    data = {field: extract_by_rule(text, rule) for field, rule in template.items()}
    return data

def extract_by_rule(text, rule):
    import re
    match = re.search(rule, text)
    return match.group(1) if match else None

# === app/email_fetcher.py ===
import imaplib
import email
import os
from dotenv import load_dotenv

load_dotenv()

def fetch_latest_email():
    mail = imaplib.IMAP4_SSL("imap.gmail.com")
    mail.login(os.getenv("EMAIL_USER"), os.getenv("EMAIL_PASS"))
    mail.select("inbox")
    _, data = mail.search(None, 'ALL')
    ids = data[0].split()
    latest_id = ids[-1]
    _, msg_data = mail.fetch(latest_id, "(RFC822)")
    raw_email = msg_data[0][1]
    msg = email.message_from_bytes(raw_email)
    for part in msg.walk():
        if part.get_content_maintype() == 'multipart':
            continue
        if part.get('Content-Disposition') is None:
            continue
        filename = part.get_filename()
        if filename:
            filepath = os.path.join("/tmp", filename)
            with open(filepath, 'wb') as f:
                f.write(part.get_payload(decode=True))
            return filepath
    return None

# === app/utils.py ===
# (Future use: helpers for template rule building, logging, etc.)

# === templates/sample_template.json ===
{
    "invoice_number": "Invoice Number: (\\d+)",
    "total_amount": "Total: \\$(\\d+\\.\\d{2})"
}